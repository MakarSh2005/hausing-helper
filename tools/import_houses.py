#!/usr/bin/env python3
"""
Конвертирует заполненный шаблон «Справочник-домов-Казани.xlsx» в
backend/prisma/reference/kazan-houses.json, который читает seed.

    python3 tools/import_houses.py путь/к/Справочник.xlsx [--verified]

--verified — УК проверены по официальному источнику: бот покажет их телефоны.
Без флага УК помечаются «данные уточняются», контакты жильцам не показываются.
Коды домов (kzn_0001…) стабильны: для уже известного адреса код сохраняется.
"""
import json, re, sys, pathlib
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / 'backend' / 'prisma' / 'reference' / 'kazan-houses.json'

def inn_ok(inn: str) -> bool:
    if not re.fullmatch(r'\d{10}', inn): return False
    w = [2, 4, 10, 3, 5, 9, 4, 6, 8]
    return (sum(int(inn[i]) * w[i] for i in range(9)) % 11) % 10 == int(inn[9])

def s(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()

def inn_of(v):
    # Из Excel ИНН приходит числом или текстом с невидимыми символами (например, U+202D из копирования
    # с сайта) — оставляем только цифры.
    return re.sub(r'\D', '', s(v))

# Как houseKey в backend/src/domain/address.ts: тип улицы не важен, регистр и точки тоже
STREET_TYPES = {'ул', 'улица', 'пр', 'пр-т', 'пркт', 'проспект', 'пер', 'переулок', 'б-р', 'бульвар', 'ш', 'шоссе',
                'пл', 'площадь', 'наб', 'набережная', 'проезд', 'тракт', 'им', 'имени', 'пр-кт', 'тер', 'территория',
                'мкр', 'микрорайон'}
def norm_key(street, number, building):
    n = lambda x: re.sub(r'\s+', ' ', re.sub(r'[«»"\'.,;:()]', ' ', x.lower().replace('ё', 'е'))).strip()
    st = ' '.join(t for t in n(street).split(' ') if t and t not in STREET_TYPES)
    return (st, n(number).replace(' ', ''), n(building).replace(' ', ''))

def num(v):
    try: return int(float(v)) if s(v) else None
    except ValueError: return None

def main():
    if len(sys.argv) < 2: sys.exit(__doc__)
    verified = '--verified' in sys.argv
    wb = load_workbook(sys.argv[1], data_only=True)
    # Прежний справочник: по адресу берём код дома (диплинки не ломаются) и недостающие сведения —
    # новая выгрузка может не знать, например, число подъездов. Регистр букв в номере не важен: 26а = 26А.
    old = {}
    if OUT.exists():
        for h in json.loads(OUT.read_text('utf-8'))['houses']:
            old[(h['street'], h['houseNumber'].lower(), h['building'].lower())] = h

    orgs, problems = [], []
    for r in list(wb['УК'].iter_rows(values_only=True))[1:]:
        if not s(r[0]) or 'ПРИМЕР' in s(r[9]).upper(): continue
        inn = inn_of(r[0])
        if not inn_ok(inn): problems.append(f'УК «{s(r[1])}»: ИНН {inn} не проходит проверку контрольной цифры')
        orgs.append(dict(inn=inn, name=s(r[1]), phone=s(r[2]), dispatcherPhone=s(r[3]), address=s(r[4]),
                         email=s(r[5]), website=s(r[6]), workingHours=s(r[7]), licenseNumber=s(r[8]),
                         note=s(r[9]), innValid=inn_ok(inn),
                         # «не подтверждён» в примечании — эта УК остаётся «данные уточняются» даже с --verified
                         verified=verified and 'не подтвержд' not in s(r[9]).lower()))
    inns = {o['inn'] for o in orgs}

    houses, used, seen = [], {h['code'] for h in old.values()}, {}
    nxt = max([int(c.split('_')[1]) for c in used] + [0]) + 1
    for i, r in enumerate(list(wb['Дома'].iter_rows(values_only=True))[1:], start=2):
        if not s(r[0]) or 'ПРИМЕР' in s(r[10]).upper(): continue
        street, number, building = re.sub(r'^тер\.\.', 'тер.', s(r[0])), s(r[1]), s(r[2])
        if not number:
            problems.append(f'Дома, строка {i}: у дома на «{street}» не указан номер — строка пропущена')
            continue
        # «15 корпус 1», «29б корп. 2», «15к1» в столбце «Дом» — корпус переносим в свой столбец
        m = re.fullmatch(r'(.+?)\s*(?:корпус|корп\.?|к\.?)\s*(\d{1,3})', number, re.IGNORECASE)
        if m and not building:
            problems.append(f'Дома, строка {i}: «{number}» — корпус перенесён в столбец «Корпус»')
            number, building = m.group(1).strip(), m.group(2)
        # «3 корпус а» — буква вместо номера корпуса: это литера дома, «3А»
        m = re.fullmatch(r'(\d+)\s*(?:корпус|корп\.?|к\.?)\s*([а-яё])', number, re.IGNORECASE)
        if m:
            problems.append(f'Дома, строка {i}: «{number}» записан как «{m.group(1)}{m.group(2).upper()}»')
            number = m.group(1) + m.group(2).upper()
        key = (street, number, building)
        nk = norm_key(street, number, building)
        if nk in seen:
            problems.append(f'Дома, строка {i}: «{street} {number}» повторяет строку {seen[nk]} — пропущена')
            continue
        seen[nk] = i
        if inn_of(r[4]) not in inns: problems.append(f'Дома, строка {i}: УК с ИНН {inn_of(r[4])} нет на листе «УК»')
        prev = old.get((key[0], key[1].lower(), key[2].lower()), {})
        code = prev.get('code')
        if not code:
            code = f'kzn_{nxt:04d}'; nxt += 1
        h = dict(code=code, street=key[0], houseNumber=key[1], building=key[2], district=s(r[3]),
                 managerInn=inn_of(r[4]), yearBuilt=num(r[5]), floors=num(r[6]), entrances=num(r[7]),
                 apartmentsCount=num(r[8]), gisUrl=s(r[9]), note=s(r[10]))
        for f in ('district', 'yearBuilt', 'floors', 'entrances', 'apartmentsCount'):
            if h[f] in (None, '') and prev.get(f) not in (None, ''): h[f] = prev[f]
        houses.append(h)

    info = {s(r[0]).rstrip('*'): s(r[1]) for r in wb['Сведения'].iter_rows(values_only=True) if r and s(r[0])}
    # Дата среза: из листа «Сведения» или ключом --date ДД.ММ.ГГГГ
    date_arg = next((a.split('=', 1)[1] for a in sys.argv if a.startswith('--date=')), '')
    actual = date_arg or info.get('Дата, на которую собраны данные', '')
    if hasattr(actual, 'strftime'): actual = actual.strftime('%d.%m.%Y')
    iso = re.fullmatch(r'(\d{4})-(\d{2})-(\d{2})', str(actual))
    if iso: actual = f'{iso.group(3)}.{iso.group(2)}.{iso.group(1)}'
    if not re.fullmatch(r'\d{2}\.\d{2}\.\d{4}', str(actual)):
        problems.append('лист «Сведения»: не указана дата среза (ДД.ММ.ГГГГ) — укажите её или запустите с --date=ДД.ММ.ГГГГ')
    data = dict(city='Казань', verified=verified, source=info.get('Источник', ''), actualAt=str(actual),
                collectedBy=info.get('Кто собирал', ''), organizations=orgs, houses=houses)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', 'utf-8')
    print(f'Записано: {len(houses)} домов, {len(orgs)} УК → {OUT.relative_to(ROOT)}')
    for p in problems: print('ВНИМАНИЕ:', p)

main()
